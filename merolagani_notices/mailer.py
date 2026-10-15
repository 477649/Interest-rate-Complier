"""Email the Excel report with a short summary.

    python -m merolagani_notices.mailer --report reports/Interest_Rate_Summary.xlsx

Sends at most two emails per month:
  1. when at least --min-coverage (default 90%) of report banks have published
     (this email lists how many notices are still pending), and
  2. again when 100% of report banks have published (marked as the complete report).
If coverage goes straight to 100%, only one email is sent.

SMTP settings come from environment variables (GitHub secrets in the workflow):
    SMTP_HOST, SMTP_PORT (587 = STARTTLS, 465 = SSL), SMTP_USERNAME, SMTP_PASSWORD,
    MAIL_TO (comma-separated), MAIL_CC (optional), MAIL_FROM (optional, defaults to SMTP_USERNAME)
"""

from __future__ import annotations

import argparse
import html
import json
import logging
import os
import smtplib
import ssl
import sys
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger("merolagani_notices.mailer")

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COMPLETE_NOTE = "Update: all banks have now published their notices, so this is the complete report."


def summarize(report: Path) -> dict:
    """Read the comparison sheet: title, and per-bank counts of increased/decreased rates."""
    ws = load_workbook(report, data_only=True)["Interest Rate Summary"]
    title = ws["A1"].value or "Interest Rate Summary"
    change_cols = [c for c in range(2, ws.max_column + 1)
                   if str(ws.cell(row=4, column=c).value or "").endswith("Changes")]
    banks, sector = [], ""
    for r in range(5, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name or str(name).startswith(("Changes =", "Rule")):
            continue
        values = [ws.cell(row=r, column=c).value for c in change_cols]
        if all(v is None for v in values):  # sector divider row
            sector = str(name)
            continue
        up = sum(isinstance(v, (int, float)) and v > 0 for v in values)
        down = sum(isinstance(v, (int, float)) and v < 0 for v in values)
        banks.append({"bank": name, "sector": sector, "up": up, "down": down})
    return {"title": title, "banks": banks,
            "up": sum(b["up"] for b in banks), "down": sum(b["down"] for b in banks)}


def pending_note(cov: dict | None) -> str:
    """Text for a partial-coverage email: how many published, how many and which are pending."""
    if not cov or not cov.get("total") or not cov.get("missing"):
        return ""
    left = len(cov["missing"])
    return (f"Note: {cov['found']} of {cov['total']} banks have published so far. "
            f"{left} notice{'s' if left != 1 else ''} pending: {', '.join(cov['missing'])}. "
            "An updated report will be emailed once all banks have published.")


def build_message(report: Path, sender: str, to: list[str], cc: list[str],
                  all_published: bool = False, cov: dict | None = None) -> EmailMessage:
    s = summarize(report)
    changed = [b for b in s["banks"] if b["up"] or b["down"]]
    subject = s["title"].replace("Bank Deposit Interest Rates — ", "Deposit Interest Rate Report: ")
    if all_published:
        subject += " (Complete: all banks published)"

    note = COMPLETE_NOTE if all_published else pending_note(cov)

    lines = [f"{b['bank']}: {b['up']} up, {b['down']} down" for b in changed]
    text = (f"Dear Sir/Madam,\n\nPlease find attached the {s['title']}.\n\n"
            + (note + "\n\n" if note else "")
            + f"Banks covered: {len(s['banks'])}\nRates increased: {s['up']}\nRates decreased: {s['down']}\n\n"
            + ("Banks with rate changes:\n" + "\n".join(lines) if lines else "No rate changes from the previous month.")
            + "\n\nThis report was generated automatically from bank notices published on merolagani.com.\n")

    rows = "".join(
        f"<tr><td style='padding:4px 10px;border:1px solid #C9D3E0'>{html.escape(str(b['bank']))}</td>"
        f"<td style='padding:4px 10px;border:1px solid #C9D3E0;text-align:center;color:#3F6F12'>{b['up']}</td>"
        f"<td style='padding:4px 10px;border:1px solid #C9D3E0;text-align:center;color:#B42318'>{b['down']}</td></tr>"
        for b in changed)
    table = (f"<table style='border-collapse:collapse;font-size:13px'><tr style='background:#0B4EA2;color:#fff'>"
             f"<th style='padding:6px 10px'>Bank</th><th style='padding:6px 10px'>Rates up</th>"
             f"<th style='padding:6px 10px'>Rates down</th></tr>{rows}</table>") if rows else \
        "<p>No rate changes from the previous month.</p>"

    if all_published:
        note_html = f"<p style='color:#0B4EA2'><b>{html.escape(COMPLETE_NOTE)}</b></p>"
    elif note:
        left = len(cov["missing"])
        note_html = (
            "<p style='background:#FFF7E6;border:1px solid #F5C26B;padding:8px 12px;color:#7A4B00'>"
            f"<b>Note:</b> {cov['found']} of {cov['total']} banks have published so far. "
            f"<b>{left} notice{'s' if left != 1 else ''} pending:</b> "
            f"{html.escape(', '.join(cov['missing']))}.<br>"
            "An updated report will be emailed once all banks have published.</p>")
    else:
        note_html = ""

    body_html = f"""<div style="font-family:Calibri,Arial,sans-serif;font-size:14px;color:#1F2937">
<p>Dear Sir/Madam,</p>
<p>Please find attached the <b>{html.escape(s['title'])}</b>.</p>
{note_html}
<p><b>Banks covered:</b> {len(s['banks'])} &nbsp;|&nbsp;
<span style="color:#3F6F12"><b>Rates increased:</b> {s['up']}</span> &nbsp;|&nbsp;
<span style="color:#B42318"><b>Rates decreased:</b> {s['down']}</span></p>
{table}
<p style="color:#6B7280;font-size:12px">Generated automatically from bank notices published on merolagani.com.</p>
</div>"""

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = ", ".join(to)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg.set_content(text)
    msg.add_alternative(body_html, subtype="html")
    msg.add_attachment(report.read_bytes(), maintype=XLSX.split("/")[0], subtype=XLSX.split("/")[1],
                       filename=report.name)
    return msg


def _split(value: str | None) -> list[str]:
    return [a.strip() for a in (value or "").replace(";", ",").split(",") if a.strip()]


def send(msg: EmailMessage, host: str, port: int, username: str, password: str) -> None:
    context = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=context, timeout=60) as smtp:
            smtp.login(username, password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as smtp:
            smtp.starttls(context=context)
            smtp.login(username, password)
            smtp.send_message(msg)


def report_banks(store: dict[str, dict[str, dict]]) -> list[str]:
    """Banks shown on the Interest Rate Summary sheet (same selection as the report)."""
    from .report import SUMMARY_BANKS
    banks = []
    for symbol, bank_months in store.items():
        sector = bank_months[max(bank_months)].get("sector", "")
        if sector in SUMMARY_BANKS and symbol not in SUMMARY_BANKS[sector]:
            continue
        banks.append(symbol)
    return sorted(banks)


def coverage(history_dir: Path) -> dict:
    """How many report banks have a notice for the newest month: {month, found, total, missing}."""
    from . import history, months
    store = history.load(history_dir) if history_dir.exists() else {}
    banks = report_banks(store)
    newest = max((m for s in banks for m in store[s]), default=None)
    found = [s for s in banks if newest in store[s]]
    return {"month": newest, "label": months.label(newest) if newest else "", "found": len(found),
            "total": len(banks), "missing": [s for s in banks if s not in found]}


def load_state(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def emailed_complete(state: dict) -> bool:
    """Was the last email for this month already sent at 100% coverage?

    Works with old state files too (they have only "banks_published": "found/total").
    """
    if "complete" in state:
        return bool(state["complete"])
    found, _, total = str(state.get("banks_published", "")).partition("/")
    return found.isdigit() and total.isdigit() and int(total) > 0 and found == total


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="merolagani_notices.mailer", description="Email the Excel report.")
    p.add_argument("--report", type=Path, default=Path("reports/Interest_Rate_Summary.xlsx"))
    p.add_argument("--history", type=Path, default=Path("notices/extracted/history"),
                   help="month-wise history, used to check how many banks have published")
    p.add_argument("--state", type=Path, default=Path("notices/email_state.json"),
                   help="remembers which month was already emailed")
    p.add_argument("--min-coverage", type=float, default=0.90,
                   help="send the first email when this share of report banks has a notice for the new month "
                        "(default 0.90); a second email is sent when all banks have published")
    p.add_argument("--force", action="store_true", help="send regardless of coverage and previous sends")
    p.add_argument("--dry-run", action="store_true", help="build the email and print it, without sending")
    args = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")

    cov = coverage(args.history)
    state = load_state(args.state)
    share = cov["found"] / cov["total"] if cov["total"] else 0
    complete = cov["total"] > 0 and cov["found"] == cov["total"]
    log.info("%s: %d of %d report banks have published (%.0f%%, need %.0f%%).", cov["label"] or "No month",
             cov["found"], cov["total"], share * 100, args.min_coverage * 100)
    if cov["missing"]:
        log.info("  Still waiting for: %s", ", ".join(cov["missing"]))

    already_emailed = state.get("last_emailed_month") == cov["month"]
    # Second email: this month was emailed before reaching 100%, and now it is 100%.
    complete_update = already_emailed and complete and not emailed_complete(state)

    if not args.force and not args.dry_run:
        if share < args.min_coverage:
            log.info("Email not sent: waiting for more banks to publish.")
            return 0
        if already_emailed and not complete_update:
            if emailed_complete(state):
                log.info("Email not sent: the complete %s report was already emailed on %s.", cov["label"],
                         state.get("emailed_at", "an earlier run"))
            else:
                log.info("Email not sent: the %s report was already emailed on %s at %s; "
                         "will email again when all banks have published.", cov["label"],
                         state.get("emailed_at", "an earlier run"), state.get("banks_published", "partial coverage"))
            return 0

    env = os.environ
    to, cc = _split(env.get("MAIL_TO")), _split(env.get("MAIL_CC"))
    username = env.get("SMTP_USERNAME", "")
    sender = env.get("MAIL_FROM") or username
    if not args.report.exists():
        log.error("Report not found: %s", args.report)
        return 1
    msg = build_message(args.report, sender or "report@example.com", to or ["(not set)"], cc,
                        all_published=complete_update, cov=cov)
    if args.dry_run:
        log.info("Subject: %s\nTo: %s\n\n%s", msg["Subject"], msg["To"],
                 msg.get_body(("plain",)).get_content())
        return 0

    missing = [k for k in ("SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "MAIL_TO") if not env.get(k)]
    if missing:
        log.warning("Email not sent: missing settings %s", ", ".join(missing))
        return 0  # not an error: email is optional
    send(msg, env["SMTP_HOST"], int(env.get("SMTP_PORT") or 587), username, env["SMTP_PASSWORD"])
    log.info("Report emailed to %s%s%s", ", ".join(to), f" (cc {', '.join(cc)})" if cc else "",
             " [complete: all banks published]" if complete_update
             else (f" [{len(cov['missing'])} pending]" if cov["missing"] else ""))
    args.state.parent.mkdir(parents=True, exist_ok=True)
    args.state.write_text(json.dumps({
        "last_emailed_month": cov["month"], "month_label": cov["label"],
        "emailed_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "banks_published": f"{cov['found']}/{cov['total']}", "complete": complete,
        "forced": args.force,
    }, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
