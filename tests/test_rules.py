"""Offline tests for the extraction rules and the Excel report."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from merolagani_notices.extract import to_record
from merolagani_notices.report import build_report
from merolagani_notices.rules import NS, fd_buckets, merge_partial, saving_min_max


def row(start, end, rate, inclusive=False):
    return {"from_months": start, "to_months": end, "to_inclusive": inclusive, "rate": rate}


class SavingTests(unittest.TestCase):
    def test_example_from_rules(self):
        self.assertEqual(saving_min_max([2.75, 2.80, 3.00, 3.25, 4.25])[:2], (2.75, 3.25))

    def test_single_rate(self):
        smin, smax, note = saving_min_max([2.75, 2.75])
        self.assertEqual((smin, smax), (2.75, 2.75))
        self.assertIn("Only one", note)

    def test_empty(self):
        self.assertEqual(saving_min_max([])[:2], (NS, NS))


class FdBucketTests(unittest.TestCase):
    def test_standard_bands(self):  # 3-6M 2.80, 6-12M 2.90, 1-2Y 3.00, 2Y+ 4.00
        rows = [row(3, 6, 2.80), row(6, 12, 2.90), row(12, 24, 3.00), row(24, None, 4.00)]
        self.assertEqual(fd_buckets(rows), (2.90, 3.00, 4.00))

    def test_range_spanning_one_year(self):  # "3 months to 2 years" 2.75
        rows = [row(3, 24, 2.75), row(24, 60, 3.05), row(60, None, 3.76)]
        self.assertEqual(fd_buckets(rows), (2.75, 2.75, 3.76))

    def test_inclusive_up_to_one_year(self):  # "6 months to 1 year" 2.85, "above 1 to 2 years" 3.00
        rows = [row(3, 6, 2.80), row(6, 12, 2.85, inclusive=True), row(12, 24, 3.00)]
        self.assertEqual(fd_buckets(rows), (2.85, 3.00, 3.00))

    def test_starts_at_one_year(self):
        self.assertEqual(fd_buckets([row(12, 24, 2.75), row(24, None, 3.50)]), (NS, 2.75, 3.50))

    def test_single_open_rate(self):  # "6 months and above" 2.75
        self.assertEqual(fd_buckets([row(6, None, 2.75)]), (2.75, 2.75, 2.75))


class FdPointTests(unittest.TestCase):
    def test_points_from_bands(self):
        from merolagani_notices.rules import fd_points
        # Kamana Sewa style: 3M-2Y 2.75, 2-3Y 3.00, 3-5Y 3.15, 5Y+ 4.28
        rows = [row(3, 24, 2.75), row(24, 36, 3.00), row(36, 60, 3.15), row(60, None, 4.28)]
        self.assertEqual(fd_points(rows), {"3M": 2.75, "6M": 2.75, "1Y": 2.75, "2Y": 3.00,
                                           "3Y": 3.15, "4Y": 3.15, "5Y+": 4.28})

    def test_points_missing_short_tenures(self):
        from merolagani_notices.rules import fd_points
        pts = fd_points([row(12, 24, 2.75), row(24, None, 3.35)])
        self.assertEqual((pts["3M"], pts["6M"], pts["1Y"], pts["5Y+"]), (NS, NS, 2.75, 3.35))


class RecordAndReportTests(unittest.TestCase):
    NOTICE = {"symbol": "TEST", "company": "Test Bank Limited", "sector": "Commercial Banks",
              "announcement_id": "10", "date": "2026-09-16", "source_url": "https://example.com"}
    RAW = {"bank_name": "Test Bank Ltd.", "effective_date": "1 Ashwin 2083", "is_partial_amendment": False,
           "saving_rates": [{"product": "Normal", "rate": 2.75}, {"product": "Remit", "rate": 3.75}],
           "call": "Up to 1.37%",
           "individual_fd": [row(3, 12, 2.90) | {"tenure_text": "3M-1Y"},
                             row(12, None, 3.50) | {"tenure_text": "1Y+"}],
           "institutional_fd": [], "notes": [], "unclear": ["blurry footer"]}

    def test_to_record_applies_rules(self):
        rec = to_record(self.RAW, self.NOTICE)
        self.assertEqual((rec["ind_lt1"], rec["ind_1y"], rec["ind_gt1"]), (2.90, 3.50, 3.50))
        self.assertEqual(rec["inst_1y"], NS)
        self.assertIn("Unclear: blurry footer", rec["notes"])

    def test_partial_amendment_keeps_previous_values(self):
        new = {"announcement_id": 11, "saving_rates": [], "call": NS, "ind_1y": NS, "notes": []}
        old = {"announcement_id": 10, "saving_rates": [2.75], "call": "Up to 0.20%", "ind_1y": 2.75}
        merged = merge_partial(new, old)
        self.assertEqual((merged["call"], merged["ind_1y"], merged["saving_rates"]), ("Up to 0.20%", 2.75, [2.75]))

    def test_month_comparison_report(self):
        from merolagani_notices import history

        bhadra = to_record(self.RAW, {**self.NOTICE, "announcement_id": "9", "date": "2026-08-16"})
        bhadra["effective"] = "1 Bhadra 2083"
        bhadra["ind_1y"] = 3.25            # Ashwin has 3.50 -> +0.25
        bhadra["call"] = "Up to 1.42%"     # Ashwin 'Up to 1.37%' -> -0.05
        ashwin = to_record(self.RAW, self.NOTICE)
        with tempfile.TemporaryDirectory() as tmp:
            hist = Path(tmp) / "history"
            self.assertEqual(history.save(hist, bhadra), "2083-05")
            self.assertEqual(history.save(hist, ashwin), "2083-06")
            path = Path(tmp) / "out.xlsx"
            build_report([ashwin], path, history_dir=hist)
            wb = load_workbook(path)
            ws = wb["Interest Rate Summary"]
            # Bank | Saving: Min cur, Min prev, Changes, Max cur, Max prev, Changes | Call cur, prev, Changes | ...
            self.assertEqual(ws["B3"].value, "Saving")
            self.assertEqual((ws["B4"].value, ws["C4"].value, ws["D4"].value),
                             ("Min Ashwin 2083", "Min Bhadra 2083", "Changes"))
            self.assertEqual((ws["H4"].value, ws["I4"].value), ("Ashwin 2083", "Bhadra 2083"))
            self.assertEqual(ws["A6"].value, "Test Bank Ltd.")
            self.assertEqual((ws["H6"].value, ws["I6"].value), ("Up to 1.37%", "Up to 1.42%"))
            self.assertAlmostEqual(ws["J6"].value, -0.0005)             # call change
            self.assertAlmostEqual(ws["N6"].value, 0.035)               # Individual 1Y, Ashwin
            self.assertAlmostEqual(ws["O6"].value, 0.0325)              # Individual 1Y, Bhadra
            self.assertAlmostEqual(ws["P6"].value, 0.0025)              # +0.25 points
            self.assertEqual(ws["D6"].value, 0)                         # saving min unchanged
            self.assertTrue(ws.conditional_formatting)                  # green/red colouring present
            self.assertEqual(wb.sheetnames,
                             ["Interest Rate Summary", "Development FD Spread", "Monthly History", "Notes"])
            self.assertEqual(wb["Monthly History"].max_row, 3)          # two months stored

    def test_history_keeps_newer_notice_within_month(self):
        from merolagani_notices import history

        with tempfile.TemporaryDirectory() as tmp:
            hist = Path(tmp)
            newer = to_record(self.RAW, {**self.NOTICE, "announcement_id": "12"})
            older = to_record(self.RAW, {**self.NOTICE, "announcement_id": "11"})
            history.save(hist, newer)
            history.save(hist, older)
            self.assertEqual(history.load(hist)["TEST"]["2083-06"]["announcement_id"], 12)


if __name__ == "__main__":
    unittest.main()
